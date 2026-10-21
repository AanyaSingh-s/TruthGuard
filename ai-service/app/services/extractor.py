"""Document text extraction engine supporting TXT, PDF, and DOCX."""

import hashlib
import io
import logging
from typing import Any, Dict, Tuple
import pypdf
import docx

logger = logging.getLogger(__name__)


class ExtractionError(Exception):
    """Base class for document extraction errors."""
    pass


class EmptyDocumentError(ExtractionError):
    """Raised when a document contains no readable text."""
    pass


class MalformedDocumentError(ExtractionError):
    """Raised when a document is corrupt or structurally unreadable."""
    pass


class UnsupportedFormatError(ExtractionError):
    """Raised when an uploaded file format is not supported."""
    pass


def compute_sha256(text: str) -> str:
    """Compute standard SHA-256 hex digest of a string."""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def extract_from_txt(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extract and normalize text from plain text file."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            text = file_bytes.decode(encoding)
            clean_text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
            if not clean_text:
                raise EmptyDocumentError("The plain text document contains no readable text")
            return clean_text, {"encoding": encoding}
        except UnicodeDecodeError:
            continue
    raise MalformedDocumentError("Failed to decode text file with supported encodings (UTF-8, Latin-1, CP1252)")


def extract_from_pdf(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extract and normalize text from PDF document."""
    try:
        stream = io.BytesIO(file_bytes)
        reader = pypdf.PdfReader(stream)
        num_pages = len(reader.pages)
        if num_pages == 0:
            raise EmptyDocumentError("PDF document contains 0 pages")

        pages_text = []
        for idx, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    pages_text.append(page_text.strip())
            except Exception as e:
                logger.warning("Failed to extract page %d from PDF: %s", idx + 1, e)

        full_text = "\n\n".join(pages_text).strip()
        if not full_text:
            raise EmptyDocumentError("PDF document contains no extractable text (document may be empty or scanned images)")

        return full_text, {
            "page_count": num_pages,
            "extracted_pages": len(pages_text)
        }
    except EmptyDocumentError:
        raise
    except Exception as e:
        logger.error("Failed to parse PDF: %s", e, exc_info=True)
        raise MalformedDocumentError(f"Malformed or corrupt PDF document: {str(e)}") from e


def extract_from_docx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extract and normalize text from DOCX document."""
    try:
        stream = io.BytesIO(file_bytes)
        doc = docx.Document(stream)

        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        # Also extract table text if present
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_texts.append(" | ".join(row_cells))

        combined = "\n".join(paragraphs + table_texts).strip()
        if not combined:
            raise EmptyDocumentError("DOCX document contains no readable paragraph or table text")

        return combined, {
            "paragraph_count": len(paragraphs),
            "table_count": len(doc.tables)
        }
    except EmptyDocumentError:
        raise
    except Exception as e:
        logger.error("Failed to parse DOCX: %s", e, exc_info=True)
        raise MalformedDocumentError(f"Malformed or corrupt DOCX document: {str(e)}") from e


def extract_text(filename: str, file_bytes: bytes) -> Tuple[str, str, Dict[str, Any]]:
    """
    Dispatch file bytes to appropriate extractor based on file extension.
    Returns: (detected_format, extracted_text, metadata)
    """
    if not file_bytes or len(file_bytes.strip()) == 0:
        raise EmptyDocumentError("Uploaded file is empty (0 bytes)")

    lower_name = filename.lower()
    if lower_name.endswith(".txt"):
        text, meta = extract_from_txt(file_bytes)
        return "TXT", text, meta
    elif lower_name.endswith(".pdf"):
        text, meta = extract_from_pdf(file_bytes)
        return "PDF", text, meta
    elif lower_name.endswith(".docx"):
        text, meta = extract_from_docx(file_bytes)
        return "DOCX", text, meta
    else:
        raise UnsupportedFormatError(
            f"Unsupported file format for '{filename}'. Only .txt, .pdf, and .docx are supported."
        )
